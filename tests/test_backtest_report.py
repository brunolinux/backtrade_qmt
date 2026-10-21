from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import pandas as pd
from openpyxl import load_workbook

from backtest_report import build_backtest_tables, write_backtest_workbook
from strategy_common import Trade


class BacktestReportTest(unittest.TestCase):
    def setUp(self):
        self.time = pd.Timestamp('2026-01-06')
        self.portfolio = pd.DataFrame([
            (self.time, '159650', 100., 262.6),
            (self.time, '518880', -200., -2000.),
            (self.time, 'cash', 2737.405, 2737.405),
        ], columns=['datetime', 'symbol', 'quantity', 'value']).set_index('datetime')
        self.trades = pd.DataFrame([
            ('159650', 1, self.time, self.time, 1.005, '止盈')], columns=Trade._fields)
        self.names = {'159650': '示例基金', '518880': '黄金基金'}

    def test_tables_have_explicit_columns_and_cent_values(self):
        tables = build_backtest_tables(self.portfolio, self.trades, self.names)
        self.assertEqual(list(tables.trade.columns), ['name', *Trade._fields])
        self.assertEqual(tables.trade.iloc[0]['name'], '示例基金')
        self.assertEqual(tables.trade.iloc[0].pnl, 1.01)
        self.assertEqual(tables.portfolio_value.columns[-1], 'total_value')
        self.assertEqual(tables.portfolio_value.iloc[0]['total_value'], 1000.01)
        self.assertEqual(tables.portfolio_value.iloc[0]['cash'], 2737.41)
        self.assertNotIn('cash', tables.portfolio_qty)
        self.assertNotIn('total_value', tables.portfolio_qty)
        self.assertEqual(tables.portfolio_qty.index.name, 'datetime')

    def test_repeated_builds_leave_all_inputs_unchanged(self):
        original = deepcopy((self.portfolio, self.trades, self.names))
        first = build_backtest_tables(self.portfolio, self.trades, self.names)
        second = build_backtest_tables(self.portfolio, self.trades, self.names)
        for a, b in zip(first, second):
            pd.testing.assert_frame_equal(a, b)
        pd.testing.assert_frame_equal(self.portfolio, original[0])
        pd.testing.assert_frame_equal(self.trades, original[1])
        self.assertEqual(self.names, original[2])

    def test_empty_results_preserve_headers_without_fake_dates(self):
        tables = build_backtest_tables(self.portfolio.iloc[:0], self.trades.iloc[:0], self.names)
        self.assertTrue(all(table.empty for table in tables))
        self.assertEqual(list(tables.trade.columns), ['name', *Trade._fields])
        self.assertEqual(list(tables.portfolio_value.columns), ['cash', 'total_value'])
        self.assertEqual(tables.portfolio_qty.index.name, 'datetime')

    def test_all_cash_preserves_dates_in_quantity_table(self):
        tables = build_backtest_tables(
            self.portfolio[self.portfolio.symbol == 'cash'], self.trades.iloc[:0], self.names)
        self.assertEqual(len(tables.portfolio_qty.index), 1)
        self.assertEqual(len(tables.portfolio_qty.columns), 0)

    def test_nan_money_is_not_silently_skipped_in_total(self):
        self.portfolio.loc[self.portfolio.symbol == '159650', 'value'] = float('nan')
        tables = build_backtest_tables(self.portfolio, self.trades, self.names)
        self.assertTrue(pd.isna(tables.portfolio_value.iloc[0]['total_value']))

    def test_qmt_codes_and_signed_integer_quantities(self):
        tables = build_backtest_tables(self.portfolio, self.trades, self.names)
        self.assertEqual(tables.trade.iloc[0].symbol, '159650.SZ')
        self.assertEqual(list(tables.portfolio_qty.columns), ['159650.SZ', '518880.SH'])
        self.assertEqual(list(tables.portfolio_qty.iloc[0]), [100, -200])
        self.assertTrue(all(pd.api.types.is_integer_dtype(dtype)
                            for dtype in tables.portfolio_qty.dtypes))
        self.portfolio.loc[self.portfolio.symbol == '159650', 'quantity'] = 0.
        self.assertEqual(build_backtest_tables(
            self.portfolio, self.trades, self.names).portfolio_qty.iloc[0, 0], 0)

    def test_existing_suffix_and_name_mapping_are_idempotent(self):
        self.portfolio['symbol'] = self.portfolio.symbol.replace({'159650': '159650.SZ'})
        self.trades['symbol'] = '159650.SZ'
        tables = build_backtest_tables(self.portfolio, self.trades, self.names)
        self.assertEqual(tables.trade.iloc[0]['name'], '示例基金')
        self.assertEqual(tables.trade.iloc[0].symbol, '159650.SZ')
        self.assertIn('cash', tables.portfolio_value)
        self.assertIn('total_value', tables.portfolio_value)

    def test_invalid_quantities_are_rejected_before_integer_conversion(self):
        for quantity in [150, 100.5, float('nan'), float('inf')]:
            with self.subTest(quantity=quantity):
                self.portfolio.loc[self.portfolio.symbol == '159650', 'quantity'] = quantity
                with self.assertRaisesRegex(ValueError, '159650'):
                    build_backtest_tables(self.portfolio, self.trades, self.names)

    def test_nullable_missing_quantity_has_a_clear_symbol_error(self):
        self.portfolio['quantity'] = self.portfolio['quantity'].astype('Float64')
        self.portfolio.loc[self.portfolio.symbol == '159650', 'quantity'] = pd.NA
        with self.assertRaisesRegex(ValueError, '159650'):
            build_backtest_tables(self.portfolio, self.trades, self.names)

    def test_unknown_or_mismatched_codes_are_rejected(self):
        for symbol in ['123456', '159650.SH', '518880.SZ', '15965', 'ETF_A', '159650.sz']:
            with self.subTest(symbol=symbol):
                portfolio = self.portfolio.replace({'symbol': {'159650': symbol}})
                with self.assertRaisesRegex(ValueError, symbol):
                    build_backtest_tables(portfolio, self.trades, self.names)

    def test_normalization_collision_is_rejected(self):
        extra = self.portfolio[self.portfolio.symbol == '159650'].copy()
        extra['symbol'] = '159650.SZ'
        with self.assertRaisesRegex(ValueError, '159650.SZ'):
            build_backtest_tables(pd.concat([self.portfolio, extra]), self.trades, self.names)

    def test_workbook_has_numeric_amounts_integer_quantities_and_dates(self):
        tables = build_backtest_tables(self.portfolio, self.trades, self.names)
        original = deepcopy(tables)
        with TemporaryDirectory() as temp:
            path = Path(temp) / 'nested' / 'results.xlsx'
            write_backtest_workbook(path, tables)
            workbook = load_workbook(path)
            self.addCleanup(workbook.close)
            self.assertEqual(workbook.sheetnames, ['trade', 'portfolio_value', 'portfolio_qty'])
            trade = workbook['trade']
            self.assertEqual([cell.value for cell in trade[1]], ['name', *Trade._fields])
            self.assertEqual(trade['A2'].value, '示例基金')
            self.assertEqual(trade['B2'].value, '159650.SZ')
            self.assertEqual(trade['B2'].data_type, 's')
            self.assertEqual(trade['F2'].value, 1.01)
            self.assertEqual(trade['F2'].number_format, '0.00')
            self.assertEqual(trade['F2'].data_type, 'n')
            self.assertEqual(trade['G2'].value, '止盈')
            self.assertEqual(trade['D2'].value, self.time.to_pydatetime())
            self.assertEqual(trade['D2'].number_format, 'yyyy-mm-dd')
            self.assertEqual(workbook['portfolio_value']['D2'].value, 2737.41)
            self.assertEqual(workbook['portfolio_value']['E2'].value, 1000.01)
            self.assertEqual(workbook['portfolio_qty']['B2'].value, 100)
            self.assertEqual(workbook['portfolio_qty']['C2'].value, -200)
            self.assertEqual(workbook['portfolio_qty']['B2'].number_format, '0')
            first = {sheet.title: list(sheet.values) for sheet in workbook}
            workbook.close()
            write_backtest_workbook(path, tables)
            repeated = load_workbook(path)
            self.assertEqual({sheet.title: list(sheet.values) for sheet in repeated}, first)
            repeated.close()
        for before, after in zip(original, tables):
            pd.testing.assert_frame_equal(before, after)

    def test_empty_workbook_still_has_three_sheets_and_headers(self):
        tables = build_backtest_tables(self.portfolio.iloc[:0], self.trades.iloc[:0], self.names)
        with TemporaryDirectory() as temp:
            path = Path(temp) / 'empty.xlsx'
            write_backtest_workbook(path, tables)
            workbook = load_workbook(path)
            self.assertEqual(len(workbook.sheetnames), 3)
            self.assertEqual(list(workbook['portfolio_value'].values), [('datetime', 'cash', 'total_value')])
            self.assertEqual(list(workbook['portfolio_qty'].values), [('datetime',)])
            workbook.close()
